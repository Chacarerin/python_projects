from docx import Document

# Crear el documento
doc = Document()

# Título del CV
doc.add_heading('Currículum Vitae', 0)

# Información personal
doc.add_paragraph('Rubén Alejandro Schnettler Lucero')
doc.add_paragraph('Fecha de nacimiento: 5 de diciembre de 1984')
doc.add_paragraph('Ubicación: Viña del Mar, Chile')
doc.add_paragraph('Teléfono: +56 9 84245810 / +56 32 2701214')
doc.add_paragraph('Correo electrónico: rubenschnettlerl@gmail.com')
doc.add_paragraph('LinkedIn: https://www.linkedin.com/in/rubensch/')
doc.add_paragraph('Portfolio: http://157.173.206.185/')

# Resumen profesional
doc.add_heading('Resumen Profesional', level=1)
doc.add_paragraph(
    "Soy un profesional con experiencia en gestión y optimización de procesos, con un fuerte interés en el área de tecnología, especialmente en la programación con Python y el uso de Inteligencia Artificial (IA). He liderado proyectos enfocados en la mejora continua, la automatización de flujos de trabajo y la integración de tecnologías emergentes en la gestión empresarial. Mi enfoque es combinar habilidades analíticas con mi capacidad de gestión para generar soluciones efectivas y adaptadas a las necesi..."
)

# Experiencia laboral
doc.add_heading('Experiencia Laboral', level=1)

doc.add_paragraph('Asesor de Gestión – Universidad Técnica Federico Santa María (2021-2023)')
doc.add_paragraph(
    "Responsable de la asesoría técnica y administrativa en la gestión de recursos financieros y humanos. Coordinación de proyectos institucionales enfocados en la mejora continua, la optimización de procesos y el alineamiento con la estrategia organizacional. Implementación de estrategias de control y seguimiento de indicadores clave de desempeño (KPI)."
)
doc.add_paragraph('Asesor Rediseño Curricular – Universidad Técnica Federico Santa María (2019-2021)')
doc.add_paragraph(
    "Gestión en la implementación de un nuevo currículo académico para carreras de ingeniería, basado en competencias. Coordinación entre distintos departamentos académicos y apoyo en la evaluación continua del proceso. Desarrollo de materiales educativos que incluyen análisis de resultados y mejoras en la formación."
)

doc.add_paragraph('Coordinador de Talento, Diversidad e Inclusión – Universidad Técnica Federico Santa María (2016-2019)')
doc.add_paragraph(
    "Liderazgo de iniciativas para atraer talento y fomentar la inclusión dentro de la universidad. Coordinación de seminarios, congresos y talleres con expertos internacionales en diversidad, impulsando programas de integración y educación basada en el enfoque STEM (Ciencia, Tecnología, Ingeniería y Matemáticas)."
)

# Educación
doc.add_heading('Formación Académica', level=1)

doc.add_paragraph('MBA, Magíster en Gestión Empresarial – Universidad Técnica Federico Santa María (2020-2022)')
doc.add_paragraph(
    "Desarrollo de habilidades en la gestión de empresas, con un enfoque en la optimización de recursos, liderazgo de equipos multidisciplinarios y planificación estratégica."
)
doc.add_paragraph('Ingeniería Civil Industrial – Universidad Técnica Federico Santa María (2016-2019)')
doc.add_paragraph(
    "Enfoque en la optimización de procesos industriales, la mejora continua y el análisis de sistemas. Implementación de proyectos tecnológicos orientados a la automatización y gestión empresarial."
)

doc.add_paragraph('Ingeniería en Acuicultura – Universidad de Los Lagos (2010-2014)')
doc.add_paragraph(
    "Enfoque en investigación aplicada, análisis de datos y desarrollo de soluciones innovadoras dentro de la industria acuícola. Utilización de tecnología para la optimización de procesos en acuicultura."
)

# Cursos y Certificaciones
doc.add_heading('Cursos y Certificaciones', level=1)

doc.add_paragraph('Bootcamp en Python – Desarrollo de Software y Automatización (2023)')
doc.add_paragraph(
    "Formación intensiva en programación con Python, incluyendo desarrollo web, automatización de procesos y análisis de datos. Desarrollo de programas para la mejora de la eficiencia operativa y la automatización de tareas administrativas."
)

doc.add_paragraph('Certificación en Marketing Digital – Google (2021)')
doc.add_paragraph(
    "Curso enfocado en estrategias de marketing digital, incluyendo SEO, campañas publicitarias, gestión de redes sociales y análisis de datos para optimizar el rendimiento en línea."
)

# Aptitudes
doc.add_heading('Aptitudes', level=1)

doc.add_paragraph(
    "Gestión de Proyectos, Optimización de Procesos, Python, Desarrollo de Software, Automización de Procesos, Marketing Digital, Análisis de Datos, Programación con IA, PostgreSQL, Gestión de Equipos, Desarrollo Web."
)

# Lenguas
doc.add_heading('Idiomas', level=1)
doc.add_paragraph('Español: Nativo')
doc.add_paragraph('Inglés: Intermedio (Conversacional)')

# Guardar documento
cv_path = "/Users/alejandro/python_projects/CV_harvard/CV_Ruben_Schnettler_Harvard_Format.docx"
doc.save(cv_path)

cv_path